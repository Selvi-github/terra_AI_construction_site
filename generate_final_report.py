# generate_final_report.py
# Generates comprehensive publication-grade Word (.docx) project report for TerraAI

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from report_content import SECTIONS_CONTENT

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_docx_report():
    doc = docx.Document()
    
    # Page setup - Margins 1 inch
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    # ─────────────────────────────────────────────────────────────
    # COVER PAGE
    # ─────────────────────────────────────────────────────────────
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(8)
    run_t = p_title.add_run("TERRA·AI")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(28)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(14, 116, 144)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Autonomous Geotechnical & Environmental Intelligence Platform\nfor Construction Site Selection & Academic EIA")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(15)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(30, 41, 59)

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.paragraph_format.space_after = Pt(24)
    run_desc = p_desc.add_run("An AI-Driven Multi-Source Geospatial Decision Support System Integrating\n5 Environmental Impact Assessment Methodologies, Stacking ML Regressors, and 3D Hazard Simulation")
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(100, 116, 139)

    # Award Banner Box
    tbl_award = doc.add_table(rows=1, cols=1)
    tbl_award.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_award = tbl_award.cell(0, 0)
    set_cell_background(c_award, "F0FDF4")
    set_cell_margins(c_award, top=140, bottom=140, left=200, right=200)
    p_aw = c_award.paragraphs[0]
    p_aw.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_aw_title = p_aw.add_run("🏆 STATE-LEVEL HACKATHON AWARD RECOGNITION\n")
    r_aw_title.bold = True
    r_aw_title.font.size = Pt(11)
    r_aw_title.font.color.rgb = RGBColor(22, 101, 52)
    r_aw_body = p_aw.add_run("Winner of SPECIAL MENTION AWARD in TANCAM's Hackathon (TNWISE 2026)\nOrganized by Tamil Nadu Centre of Excellence for Advanced Manufacturing (TANCAM),\nDassault Systèmes, TIDCO & Kumaraguru College of Technology")
    r_aw_body.font.size = Pt(9.5)
    r_aw_body.font.color.rgb = RGBColor(21, 128, 61)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(36)

    # Project Details Table
    tbl_proj = doc.add_table(rows=5, cols=2)
    tbl_proj.alignment = WD_TABLE_ALIGNMENT.CENTER
    proj_meta = [
        ("Project Details", ""),
        ("Submitted by:", "Selvi. Vaira Selvi & Team"),
        ("Institution:", "Kamaraj College of Engineering and Technology, Tamil Nadu"),
        ("Department:", "Department of Computer Science and Engineering"),
        ("Academic Year:", "2025 – 2026")
    ]
    for idx, (k, v) in enumerate(proj_meta):
        cell_k = tbl_proj.cell(idx, 0)
        cell_v = tbl_proj.cell(idx, 1)
        if idx == 0:
            set_cell_background(cell_k, "0E7490")
            set_cell_background(cell_v, "0E7490")
            p = cell_k.paragraphs[0]
            r = p.add_run(k)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        else:
            set_cell_background(cell_k, "F8FAFC")
            set_cell_background(cell_v, "FFFFFF")
            p = cell_k.paragraphs[0]
            r = p.add_run(k)
            r.bold = True
            p2 = cell_v.paragraphs[0]
            p2.add_run(v)
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # HACKATHON AWARD & GALLERY PAGE
    # ─────────────────────────────────────────────────────────────
    h_cert = doc.add_heading("TNWISE 2026 Hackathon Special Mention Award", level=1)
    h_cert.paragraph_format.space_before = Pt(12)
    h_cert.paragraph_format.space_after = Pt(12)

    p_aw_desc = doc.add_paragraph()
    p_aw_desc.add_run(
        "The TerraAI project was submitted, presented, and rigorously defended at the state-level TANCAM's Hackathon: "
        "TNWISE 2026 (Exclusive for Tamil Nadu Women in Science and Engineering), organized by the Tamil Nadu Centre "
        "of Excellence for Advanced Manufacturing (TANCAM) in association with Dassault Systèmes, TIDCO, and Kumaraguru "
        "College of Technology, Coimbatore on 12th March 2026. The project competed against top collegiate engineering teams "
        "across Tamil Nadu and received the prestigious SPECIAL MENTION AWARD for its novel integration of automated "
        "geospatial sensing, civil engineering compliance per Indian Standards (IS Codes), and machine learning stacking ensembles."
    )

    img_cert_path = "tnwise_extracted/page_4_img_2_X2.jpg"
    img_stage_path = "tnwise_extracted/page_3_img_2_X2.jpg"
    img_demo_path = "tnwise_extracted/page_2_img_2_X2.jpg"
    img_team_path = "tnwise_extracted/page_1_img_2_X2.jpg"

    if os.path.exists(img_cert_path):
        doc.add_paragraph().add_run("Figure A: Official Certificate of Appreciation — TNWISE 2026 Special Mention Award").bold = True
        doc.add_picture(img_cert_path, width=Inches(5.8))
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_page_break()

    doc.add_heading("Hackathon Evaluation, Stage Felicitation & Institutional Demonstration", level=2)
    
    if os.path.exists(img_stage_path) and os.path.exists(img_demo_path):
        tbl_imgs = doc.add_table(rows=2, cols=2)
        tbl_imgs.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        c00 = tbl_imgs.cell(0, 0)
        c00.paragraphs[0].add_run("Award Ceremony on Main Stage:\n").bold = True
        c00.paragraphs[0].add_run().add_picture(img_stage_path, width=Inches(2.8))
        
        c01 = tbl_imgs.cell(0, 1)
        c01.paragraphs[0].add_run("Live Jury Demonstration (Geotagged):\n").bold = True
        c01.paragraphs[0].add_run().add_picture(img_demo_path, width=Inches(2.8))

        c10 = tbl_imgs.cell(1, 0)
        c10.paragraphs[0].add_run("Institutional Presentation & Certificate:\n").bold = True
        c10.paragraphs[0].add_run().add_picture(img_team_path, width=Inches(2.8))
        
        c11 = tbl_imgs.cell(1, 1)
        p_eval = c11.paragraphs[0]
        p_eval.add_run("Hackathon Jury Review Highlights:\n").bold = True
        p_eval.add_run(
            "• Commended for automated multi-domain satellite ETL pipeline eliminating manual survey delays.\n"
            "• Highly rated for mathematical consistency across 5 standard EIA methodologies.\n"
            "• Praised for intuitive 3D WebGL structural hazard reaction simulations."
        )
        p_eval.paragraph_format.space_before = Pt(6)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # TABLE OF CONTENTS
    # ─────────────────────────────────────────────────────────────
    h_toc = doc.add_heading("Table of Contents", level=1)
    
    toc_items = [(f"{num}", title, f"{num+1}") for num, title, _ in SECTIONS_CONTENT]

    tbl_toc = doc.add_table(rows=len(toc_items), cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (sec_no, sec_title, p_no) in enumerate(toc_items):
        c0 = tbl_toc.cell(idx, 0)
        c1 = tbl_toc.cell(idx, 1)
        c2 = tbl_toc.cell(idx, 2)
        c0.paragraphs[0].add_run(sec_no).bold = True
        c1.paragraphs[0].add_run(sec_title)
        c2.paragraphs[0].add_run(p_no)
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_margins(c0, top=30, bottom=30, left=50, right=50)
        set_cell_margins(c1, top=30, bottom=30, left=50, right=50)
        set_cell_margins(c2, top=30, bottom=30, left=50, right=50)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # ALL 49 SECTIONS
    # ─────────────────────────────────────────────────────────────
    for sec_num, sec_title, sec_body in SECTIONS_CONTENT:
        h = doc.add_heading(f"{sec_num}. {sec_title}", level=1)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(8)
        
        for para_text in sec_body.strip().split('\n\n'):
            if para_text.strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                
                # Check for subheadings
                lines = para_text.strip().split('\n')
                if len(lines) > 1 and lines[0].strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.', '19.', '20.', '27.', '29.', '32.', '33.', '43.', '47.')):
                    r_sub = p.add_run(lines[0] + '\n')
                    r_sub.bold = True
                    r_sub.font.size = Pt(12)
                    r_sub.font.color.rgb = RGBColor(14, 116, 144)
                    body_remainder = '\n'.join(lines[1:])
                    if body_remainder.strip():
                        p.add_run(body_remainder)
                else:
                    p.add_run(para_text.strip())

    docx_path = "d:/construction_site_selection/construction_site_selection/TERRA_AI_PROJECT_REPORT.docx"
    doc.save(docx_path)
    print(f"SUCCESS: Generated Word document -> {docx_path}")
    return docx_path

if __name__ == "__main__":
    create_docx_report()
